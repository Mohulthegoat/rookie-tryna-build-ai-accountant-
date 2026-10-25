"""Generate a realistic FAKE Indian company dataset for LedgerAI development.

NEVER use real client data here. Every name, GSTIN and amount is invented.

The dataset is for the month of October 2025 and contains the six deliberate
messy cases required by the product plan:

  1. missing_invoice        - payment to a vendor with no purchase invoice
  2. duplicate_payment      - the same payment debited twice by the bank
  3. split_payment          - one invoice paid by two bank transactions
  4. small_difference       - bank and ledger differ by exactly 5 paise
  5. uncleared_cheque       - cheque recorded in the books, not yet in bank
  6. unrecorded_bank_charge - bank charge that never reached the ledger

Files produced (all amounts are rupee strings with 2 decimals in the source
files; the ground-truth JSON uses whole paise integers):

  bank_statement_oct2025.csv / .xlsx
  ledger_oct2025.csv              (Tally-style day book, two lines per voucher)
  purchase_invoices_oct2025.csv
  sales_invoices_oct2025.csv
  vendors_and_customers_oct2025.csv
  ground_truth.json

The source files deliberately look like real exports (commas, DD/MM/YYYY,
separate debit/credit columns). The ingestion layer (Phase 2) is responsible
for converting them to paise integers.
"""

from __future__ import annotations

import csv
import json
import random
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from .money import paise_to_rupees_str

__all__ = ["generate_all", "COMPANY", "OPENING_BALANCE_PAISE"]

COMPANY = {
    "name": "Sharma Trading Co. Pvt. Ltd. (FAKE)",
    "gstin": "27AABCS1234F1Z5",  # 27 = Maharashtra (fake)
    "pan": "AABCS1234F",
    "state_code": "27",
    "state": "Maharashtra",
    "financial_year": "2025-26",
    "period": "October 2025",
    "currency": "INR",
}

# Opening bank balance: Rs 1,00,000.00
OPENING_BALANCE_PAISE = 10_000_000

_PERIOD_START = date(2025, 10, 1)


def _rupees(paise: int) -> str:
    """Source files show rupees, not paise. This is intentional realism."""
    return paise_to_rupees_str(paise)


def _gst_breakup(total_paise: int, rate_percent: int, inter_state: bool) -> dict:
    """Split a GST-inclusive invoice total into taxable value + GST components.

    Deterministic integer maths. The taxable value is rounded to the nearest
    paise and any rounding remainder stays in the tax amount so that
    taxable + CGST + SGST + IGST == total, always.

    This is FAKE data generation, not tax advice. The real GST engine is
    Phase 8 and every rule there needs CA review.
    """
    taxable = int(
        (Decimal(total_paise) * 100 / Decimal(100 + rate_percent)).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )
    gst = total_paise - taxable
    if inter_state:
        cgst, sgst, igst = 0, 0, gst
    else:
        cgst = gst // 2
        sgst = gst - cgst  # odd paise goes to SGST so the sum is exact
        igst = 0
    return {
        "taxable_paise": taxable,
        "cgst_paise": cgst,
        "sgst_paise": sgst,
        "igst_paise": igst,
        "gst_paise": gst,
        "total_paise": total_paise,
        "gst_percent": rate_percent,
    }


# ---------------------------------------------------------------------------
# Master data (all fake)
# ---------------------------------------------------------------------------

PARTIES = [
    {"name": "ABC Suppliers Pvt Ltd", "gstin": "27AABCA1111F1Z1", "state_code": "27", "type": "Vendor"},
    {"name": "SAME SUPPLIER", "gstin": "27AABCS2222F1Z2", "state_code": "27", "type": "Vendor"},
    {"name": "SPLIT SUPPLIER", "gstin": "27AABCS3333F1Z3", "state_code": "27", "type": "Vendor"},
    {"name": "GANESH ENTERPRISES", "gstin": "27AABCG4444F1Z4", "state_code": "27", "type": "Vendor"},
    {"name": "Krishna Traders", "gstin": "24AABCK5555F1Z5", "state_code": "24", "type": "Vendor"},
    {"name": "Mehta Services", "gstin": "27AABCM6666F1Z6", "state_code": "27", "type": "Vendor"},
    {"name": "Uncleared Supplier", "gstin": "27AABCU7777F1Z7", "state_code": "27", "type": "Vendor"},
    {"name": "Retail Customer", "gstin": "27AABCR8888F1Z8", "state_code": "27", "type": "Customer"},
    {"name": "Verma Stores", "gstin": "27AABCV9999F1Z9", "state_code": "27", "type": "Customer"},
    {"name": "Gupta & Sons", "gstin": "27AABCG1212F1Z1", "state_code": "27", "type": "Customer"},
]


# ---------------------------------------------------------------------------
# Bank statement
# ---------------------------------------------------------------------------

def _bank_rows() -> list[dict]:
    """Build bank transactions in paise. Returns rows WITHOUT opening balance."""
    return [
        {
            "date": date(2025, 10, 1),
            "narration": "UPI/CR/512345678901/RETAIL CUSTOMER/INV SAL-2001",
            "cheque_no": "",
            "debit": 0,
            "credit": 4_500_000,
            "tag": "exact_match_receipt",
        },
        {
            "date": date(2025, 10, 3),
            "narration": "NEFT/DR/N123456789/ABC SUPPLIERS PVT LTD/INV-1001",
            "cheque_no": "",
            "debit": 1_180_000,
            "credit": 0,
            "tag": "exact_match_payment",
        },
        {
            "date": date(2025, 10, 5),
            "narration": "NEFT/DR/N223344556/GANESH ENTERPRISES/ADVANCE",
            "cheque_no": "",
            "debit": 2_360_000,
            "credit": 0,
            "tag": "missing_invoice",
        },
        {
            "date": date(2025, 10, 7),
            "narration": "UPI/CR/623456789012/RETAIL CUSTOMER/INV SAL-2002",
            "cheque_no": "",
            "debit": 0,
            "credit": 3_000_000,
            "tag": "exact_match_receipt",
        },
        {
            "date": date(2025, 10, 8),
            "narration": "NEFT/DR/N445566778/SAME SUPPLIER/INV-1003",
            "cheque_no": "",
            "debit": 1_003_000,
            "credit": 0,
            "tag": "duplicate_payment_original",
        },
        {
            "date": date(2025, 10, 8),
            "narration": "NEFT/DR/N445566779/SAME SUPPLIER/INV-1003",
            "cheque_no": "",
            "debit": 1_003_000,
            "credit": 0,
            "tag": "duplicate_payment_duplicate",
        },
        {
            "date": date(2025, 10, 9),
            "narration": "UPI/CR/734567890123/GUPTA & SONS/INV SAL-2005 ROUNDED",
            "cheque_no": "",
            "debit": 0,
            "credit": 1_500_005,  # 5 paise more than the ledger
            "tag": "small_difference",
        },
        {
            "date": date(2025, 10, 12),
            "narration": "IMPS/CR/845678901234/VERMA STORES/INV SAL-2003",
            "cheque_no": "",
            "debit": 0,
            "credit": 2_000_000,
            "tag": "exact_match_receipt",
        },
        {
            "date": date(2025, 10, 15),
            "narration": "IMPS/DR/956789012345/SPLIT SUPPLIER/PART-1",
            "cheque_no": "",
            "debit": 600_000,
            "credit": 0,
            "tag": "split_payment_part_1",
        },
        {
            "date": date(2025, 10, 16),
            "narration": "IMPS/DR/956789012346/SPLIT SUPPLIER/PART-2",
            "cheque_no": "",
            "debit": 580_000,
            "credit": 0,
            "tag": "split_payment_part_2",
        },
        {
            "date": date(2025, 10, 20),
            "narration": "BANK CHARGES/NEFT CHARGES/SMS CHARGES",
            "cheque_no": "",
            "debit": 15_000,  # Rs 150.00
            "credit": 0,
            "tag": "unrecorded_bank_charge",
        },
        {
            "date": date(2025, 10, 25),
            "narration": "UPI/CR/167890123456/VERMA STORES/INV SAL-2004",
            "cheque_no": "",
            "debit": 0,
            "credit": 1_200_000,
            "tag": "exact_match_receipt",
        },
    ]


def _with_running_balance(rows: list[dict]) -> list[dict]:
    """Add the running balance a real bank statement shows."""
    balance = OPENING_BALANCE_PAISE
    out = []
    for row in rows:
        balance += row["credit"] - row["debit"]
        enriched = dict(row)
        enriched["balance"] = balance
        out.append(enriched)
    return out


# ---------------------------------------------------------------------------
# Ledger (Tally-style day book)
# ---------------------------------------------------------------------------

def _ledger_lines() -> list[dict]:
    """Two lines per voucher, so total debits always equal total credits."""

    def voucher(vno, vtype, day, narration, dr_account, cr_account, amount, cheque=""):
        return [
            {
                "date": day,
                "voucher_no": vno,
                "voucher_type": vtype,
                "ledger": dr_account,
                "debit": amount,
                "credit": 0,
                "narration": narration,
                "cheque_no": cheque,
            },
            {
                "date": day,
                "voucher_no": vno,
                "voucher_type": vtype,
                "ledger": cr_account,
                "debit": 0,
                "credit": amount,
                "narration": narration,
                "cheque_no": cheque,
            },
        ]

    lines: list[dict] = []
    lines += voucher("RV-001", "Receipt", date(2025, 10, 1), "Receipt from Retail Customer against SAL-2001", "Bank", "Sales", 4_500_000)
    lines += voucher("PV-001", "Payment", date(2025, 10, 3), "Payment to ABC Suppliers against INV-1001", "Purchases", "Bank", 1_180_000)
    # No invoice exists for this one - the deliberate missing-invoice case.
    lines += voucher("PV-002", "Payment", date(2025, 10, 5), "Advance to Ganesh Enterprises - invoice not received", "Purchases", "Bank", 2_360_000)
    lines += voucher("RV-002", "Receipt", date(2025, 10, 7), "Receipt from Retail Customer against SAL-2002", "Bank", "Sales", 3_000_000)
    lines += voucher("PV-003", "Payment", date(2025, 10, 8), "Payment to Same Supplier against INV-1003", "Purchases", "Bank", 1_003_000)
    # Ledger has Rs 15,000.00; bank shows Rs 15,000.05 - the 5 paise case.
    lines += voucher("RV-003", "Receipt", date(2025, 10, 9), "Receipt from Gupta & Sons against SAL-2005", "Bank", "Sales", 1_500_000)
    lines += voucher("RV-004", "Receipt", date(2025, 10, 12), "Receipt from Verma Stores against SAL-2003", "Bank", "Sales", 2_000_000)
    lines += voucher("PV-004", "Payment", date(2025, 10, 15), "Payment to Split Supplier - part 1 of INV-1004", "Purchases", "Bank", 600_000)
    lines += voucher("PV-005", "Payment", date(2025, 10, 16), "Payment to Split Supplier - part 2 of INV-1004", "Purchases", "Bank", 580_000)
    lines += voucher("RV-005", "Receipt", date(2025, 10, 25), "Receipt from Verma Stores against SAL-2004", "Bank", "Sales", 1_200_000)
    # Uncleared cheque: in the books, not in the bank statement.
    lines += voucher(
        "PV-006",
        "Payment",
        date(2025, 10, 26),
        "Cheque issued to Uncleared Supplier - not yet presented",
        "Purchases",
        "Bank",
        1_200_000,
        cheque="100030",
    )
    return lines


# ---------------------------------------------------------------------------
# Invoices
# ---------------------------------------------------------------------------

def _purchase_invoices() -> list[dict]:
    specs = [
        # invoice no, date, vendor, total paise, rate, inter-state
        ("INV-2025-1001", date(2025, 10, 3), "ABC Suppliers Pvt Ltd", 1_180_000, 18, False),
        ("INV-2025-1002", date(2025, 10, 6), "ABC Suppliers Pvt Ltd", 590_000, 18, False),
        ("INV-2025-1003", date(2025, 10, 7), "SAME SUPPLIER", 1_003_000, 18, False),
        ("INV-2025-1004", date(2025, 10, 14), "SPLIT SUPPLIER", 1_180_000, 18, False),
        ("INV-2025-1005", date(2025, 10, 18), "Krishna Traders", 472_000, 18, True),
        ("INV-2025-1006", date(2025, 10, 22), "Mehta Services", 236_000, 18, False),
    ]
    party_by_name = {p["name"]: p for p in PARTIES}
    out = []
    for number, day, vendor, total, rate, inter_state in specs:
        party = party_by_name[vendor]
        split = _gst_breakup(total, rate, inter_state)
        out.append(
            {
                "invoice_no": number,
                "date": day,
                "name": vendor,
                "gstin": party["gstin"],
                "state_code": party["state_code"],
                **split,
            }
        )
    return out


def _sales_invoices() -> list[dict]:
    specs = [
        ("SAL-2025-2001", date(2025, 10, 1), "Retail Customer", 4_500_000),
        ("SAL-2025-2002", date(2025, 10, 7), "Retail Customer", 3_000_000),
        ("SAL-2025-2003", date(2025, 10, 12), "Verma Stores", 2_000_000),
        ("SAL-2025-2004", date(2025, 10, 25), "Verma Stores", 1_200_000),
        ("SAL-2025-2005", date(2025, 10, 9), "Gupta & Sons", 1_500_000),
        ("SAL-2025-2006", date(2025, 10, 27), "Gupta & Sons", 708_000),
    ]
    party_by_name = {p["name"]: p for p in PARTIES}
    out = []
    for number, day, customer, total in specs:
        party = party_by_name[customer]
        split = _gst_breakup(total, 18, False)
        out.append(
            {
                "invoice_no": number,
                "date": day,
                "name": customer,
                "gstin": party["gstin"],
                "state_code": party["state_code"],
                **split,
            }
        )
    return out


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    # utf-8-sig makes Excel open Indian company names correctly.
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_bank_csv(path: Path, rows: list[dict]) -> None:
    opening = {
        "Date": _PERIOD_START.strftime("%d/%m/%Y"),
        "ValueDate": _PERIOD_START.strftime("%d/%m/%Y"),
        "Narration": "OPENING BALANCE",
        "ChequeNo": "",
        "Debit": "",
        "Credit": "",
        "Balance": _rupees(OPENING_BALANCE_PAISE),
    }
    out = [opening]
    for row in rows:
        out.append(
            {
                "Date": row["date"].strftime("%d/%m/%Y"),
                "ValueDate": row["date"].strftime("%d/%m/%Y"),
                "Narration": row["narration"],
                "ChequeNo": row["cheque_no"],
                "Debit": _rupees(row["debit"]) if row["debit"] else "",
                "Credit": _rupees(row["credit"]) if row["credit"] else "",
                "Balance": _rupees(row["balance"]),
            }
        )
    _write_csv(path, ["Date", "ValueDate", "Narration", "ChequeNo", "Debit", "Credit", "Balance"], out)


def _write_bank_xlsx(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    headers = ["Date", "ValueDate", "Narration", "ChequeNo", "Debit", "Credit", "Balance"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    opening_row = [_PERIOD_START, _PERIOD_START, "OPENING BALANCE", None, None, None, OPENING_BALANCE_PAISE / 100]
    ws.append(opening_row)
    for row in rows:
        ws.append(
            [
                row["date"],
                row["date"],
                row["narration"],
                row["cheque_no"] or None,
                (row["debit"] / 100) if row["debit"] else None,
                (row["credit"] / 100) if row["credit"] else None,
                row["balance"] / 100,
            ]
        )

    for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=7):
        for cell in row_cells:
            if cell.column in (1, 2):
                cell.number_format = "DD/MM/YYYY"
            elif cell.column in (5, 6, 7):
                cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 12
    for col in ("E", "F", "G"):
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A2"
    wb.save(path)


def _write_ledger(path: Path, rows: list[dict]) -> None:
    out = []
    for row in rows:
        out.append(
            {
                "Date": row["date"].strftime("%d/%m/%Y"),
                "VoucherNo": row["voucher_no"],
                "VoucherType": row["voucher_type"],
                "LedgerName": row["ledger"],
                "Debit": _rupees(row["debit"]) if row["debit"] else "",
                "Credit": _rupees(row["credit"]) if row["credit"] else "",
                "Narration": row["narration"],
                "ChequeNo": row["cheque_no"],
            }
        )
    _write_csv(path, ["Date", "VoucherNo", "VoucherType", "LedgerName", "Debit", "Credit", "Narration", "ChequeNo"], out)


def _write_invoices(path: Path, rows: list[dict], name_header: str) -> None:
    out = []
    for row in rows:
        out.append(
            {
                "InvoiceNo": row["invoice_no"],
                "InvoiceDate": row["date"].strftime("%d/%m/%Y"),
                name_header: row["name"],
                "PartyGSTIN": row["gstin"],
                "SupplyStateCode": row["state_code"],
                "TaxableValue": _rupees(row["taxable_paise"]),
                "GSTPercent": row["gst_percent"],
                "CGST": _rupees(row["cgst_paise"]),
                "SGST": _rupees(row["sgst_paise"]),
                "IGST": _rupees(row["igst_paise"]),
                "InvoiceTotal": _rupees(row["total_paise"]),
            }
        )
    _write_csv(
        path,
        [
            "InvoiceNo",
            "InvoiceDate",
            name_header,
            "PartyGSTIN",
            "SupplyStateCode",
            "TaxableValue",
            "GSTPercent",
            "CGST",
            "SGST",
            "IGST",
            "InvoiceTotal",
        ],
        out,
    )


def _write_parties(path: Path) -> None:
    out = [
        {
            "Name": p["name"],
            "GSTIN": p["gstin"],
            "StateCode": p["state_code"],
            "Type": p["type"],
        }
        for p in PARTIES
    ]
    _write_csv(path, ["Name", "GSTIN", "StateCode", "Type"], out)


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def generate_all(outdir: str | Path, seed: int = 42) -> dict:
    """Generate every fake file into outdir and return the ground-truth dict."""
    random.seed(seed)
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    bank_rows = _with_running_balance(_bank_rows())
    ledger_rows = _ledger_lines()
    purchases = _purchase_invoices()
    sales = _sales_invoices()

    _write_bank_csv(outdir / "bank_statement_oct2025.csv", bank_rows)
    _write_bank_xlsx(outdir / "bank_statement_oct2025.xlsx", bank_rows)
    _write_ledger(outdir / "ledger_oct2025.csv", ledger_rows)
    _write_invoices(outdir / "purchase_invoices_oct2025.csv", purchases, "VendorName")
    _write_invoices(outdir / "sales_invoices_oct2025.csv", sales, "CustomerName")
    _write_parties(outdir / "vendors_and_customers_oct2025.csv")

    closing_balance = bank_rows[-1]["balance"]

    # Row numbers below are data rows in the CSV *excluding* the header and
    # the opening-balance row, so row 1 is the first real transaction.
    duplicate_original = next(i for i, r in enumerate(bank_rows, 1) if r["tag"] == "duplicate_payment_original")
    duplicate_second = next(i for i, r in enumerate(bank_rows, 1) if r["tag"] == "duplicate_payment_duplicate")

    truth = {
        "note": "ALL DATA IS FAKE. Amounts are whole paise integers.",
        "generated_with_seed": seed,
        "company": COMPANY,
        "period": {"start": "2025-10-01", "end": "2025-10-31"},
        "opening_balance_paise": OPENING_BALANCE_PAISE,
        "closing_balance_paise": closing_balance,
        "counts": {
            "bank_transactions": len(bank_rows),
            "ledger_lines": len(ledger_rows),
            "purchase_invoices": len(purchases),
            "sales_invoices": len(sales),
        },
        "cases": [
            {
                "type": "missing_invoice",
                "description": "Bank payment to Ganesh Enterprises with no purchase invoice on record",
                "bank_row": 3,
                "date": "2025-10-05",
                "amount_paise": 2_360_000,
                "vendor": "GANESH ENTERPRISES",
            },
            {
                "type": "duplicate_payment",
                "description": "The same payment to Same Supplier appears twice on the same day",
                "bank_rows": [duplicate_original, duplicate_second],
                "date": "2025-10-08",
                "amount_paise": 1_003_000,
                "vendor": "SAME SUPPLIER",
                "valid_invoice": "INV-2025-1003",
            },
            {
                "type": "split_payment",
                "description": "One invoice paid by two bank transactions",
                "invoice_no": "INV-2025-1004",
                "invoice_total_paise": 1_180_000,
                "bank_rows": [9, 10],
                "parts_paise": [600_000, 580_000],
            },
            {
                "type": "small_difference",
                "description": "Bank credit is 5 paise more than the ledger receipt",
                "bank_row": 7,
                "date": "2025-10-09",
                "bank_amount_paise": 1_500_005,
                "ledger_amount_paise": 1_500_000,
                "difference_paise": 5,
                "party": "Gupta & Sons",
            },
            {
                "type": "uncleared_cheque",
                "description": "Cheque 100030 is in the ledger but has not appeared in the bank statement",
                "voucher_no": "PV-006",
                "cheque_no": "100030",
                "date": "2025-10-26",
                "amount_paise": 1_200_000,
            },
            {
                "type": "unrecorded_bank_charge",
                "description": "Bank debited charges that are not in the ledger",
                "bank_row": 11,
                "date": "2025-10-20",
                "amount_paise": 15_000,
            },
        ],
        "files": {
            "bank_csv": "bank_statement_oct2025.csv",
            "bank_xlsx": "bank_statement_oct2025.xlsx",
            "ledger": "ledger_oct2025.csv",
            "purchase_invoices": "purchase_invoices_oct2025.csv",
            "sales_invoices": "sales_invoices_oct2025.csv",
            "parties": "vendors_and_customers_oct2025.csv",
        },
    }

    with open(outdir / "ground_truth.json", "w", encoding="utf-8") as handle:
        json.dump(truth, handle, indent=2, ensure_ascii=False)
        handle.write("\n")

    return truth
